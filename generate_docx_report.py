import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    PRIMARY = RGBColor(22, 101, 52)    # Deep Green #166534
    TEAL = RGBColor(15, 118, 110)      # Muted Teal #0f766e
    DARK = RGBColor(30, 41, 59)        # Slate 800 #1e293b
    MUTED = RGBColor(100, 116, 139)    # Slate 500 #64748b

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = DARK

    # Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("WattWise AI — Project Evaluation & Technical Report")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Smart Energy Consumption Optimizer for InnoTech Hackathon")
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = TEAL
    run_sub.font.bold = True

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = PRIMARY
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = TEAL
        return h

    # Section 1: Executive Summary
    add_heading_1("1. Executive Summary")
    
    table1 = doc.add_table(rows=6, cols=2)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    data_summary = [
        ("Project Name", "WattWise AI — Smart Energy Consumption Optimizer"),
        ("Hackathon", "InnoTech Hackathon"),
        ("Category", "Smart Solution & Smart Society"),
        ("Domain", "Renewable Energy & Sustainable Engineering"),
        ("Target Users", "Academic campuses, commercial facilities, SMEs, and residential users."),
        ("Status & Delivery", "Fully operational and live — React 18 / Tailwind CSS frontend actively communicating with Python 3.11 / Flask / Scikit-learn backend.")
    ]

    for idx, (k, v) in enumerate(data_summary):
        row = table1.rows[idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(1.8)
        cell_v.width = Inches(5.0)

        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(2)
        pk.paragraph_format.space_after = Pt(2)
        rk = pk.add_run(k)
        rk.font.bold = True
        rk.font.color.rgb = PRIMARY

        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(2)
        pv.paragraph_format.space_after = Pt(2)
        pv.add_run(v)

        bg = "F1F5F9" if idx % 2 == 0 else "FFFFFF"
        set_cell_background(cell_k, "E2E8F0" if idx % 2 == 0 else "F8FAFC")
        set_cell_background(cell_v, bg)
        set_cell_margins(cell_k, 50, 50, 80, 80)
        set_cell_margins(cell_v, 50, 50, 80, 80)

    # Section 2: Problem Statement & Objectives
    add_heading_1("2. Problem Statement & Key Objectives")
    doc.add_paragraph(
        "Commercial and campus facilities face unmonitored power wastage, high off-peak idle standby loads, "
        "and rising electricity bills due to lack of visibility. Most existing commercial solutions rely on expensive "
        "IoT hardware installations or act as unexplainable black boxes."
    )
    doc.add_paragraph(
        "WattWise AI solves this through an accessible, software-driven energy intelligence platform providing: "
        "(1) centralized consumption telemetry, (2) transparent rule-based energy-saving recommendations with concrete ₹ and kWh savings, "
        "(3) rigorous Scikit-learn forecasting evaluated against a chronological hold-out baseline, "
        "(4) configurable tariff and carbon footprint calculations, and (5) software-based rooftop solar PV sizing."
    )

    # Section 3: System Architecture
    add_heading_1("3. System Architecture & Technology Stack")
    
    table_tech = doc.add_table(rows=8, cols=3)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER

    tech_headers = ["Layer", "Technology", "Role & Engineering Responsibility"]
    hdr_row = table_tech.rows[0]
    for c_idx, h_text in enumerate(tech_headers):
        c = hdr_row.cells[c_idx]
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(c, "166534")
        set_cell_margins(c, 60, 60, 80, 80)

    tech_rows = [
        ("Frontend Framework", "React 18.3 + Vite 5.4", "Single-page responsive application with component isolation and rapid hot-module replacement."),
        ("Styling System", "Tailwind CSS v4", "Curated sustainability design system using deep green (#166534) and soft slate neutrals (#f7f8f6)."),
        ("Data Visualization", "Recharts 2.12", "Interactive bar, line, and composed time-series charts for daily, weekly, and forecast curves."),
        ("Iconography", "Lucide React", "Lightweight, accessible icon set (Zap, Leaf, IndianRupee, Sun, etc.)."),
        ("Backend Framework", "Python 3.11 + Flask 3.0", "RESTful API server handling CORS, file streaming, validation, calculations, and data endpoints."),
        ("Data Manipulation", "Pandas 2.2 + NumPy 1.26", "Time-series daily resampling, rolling-window calculations, and standard deviation anomaly math."),
        ("Machine Learning", "Scikit-learn 1.4", "Time-series feature engineering, chronological train/test split, and Ridge regression model.")
    ]

    for idx, (layer, tech, role) in enumerate(tech_rows):
        row = table_tech.rows[idx + 1]
        for c_idx, val in enumerate([layer, tech, role]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            if c_idx < 2:
                r.font.bold = True
            bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 40, 40, 60, 60)

    # Section 4: Feature-Verification Table
    add_heading_1("4. Feature-Verification & Code Audit Matrix")
    
    table_feat = doc.add_table(rows=18, cols=4)
    table_feat.alignment = WD_TABLE_ALIGNMENT.CENTER

    feat_headers = ["Feature / Module", "Status", "Code Evidence", "Notes & Report Adjustments"]
    f_hdr = table_feat.rows[0]
    for c_idx, h_text in enumerate(feat_headers):
        c = f_hdr.cells[c_idx]
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(c, "166534")
        set_cell_margins(c, 60, 60, 80, 80)

    verification_data = [
        ("Dashboard KPIs", "Verified Working", "calc.py:57-72, pages.jsx:17-38", "Dynamically computes total kWh, bill in ₹, emissions in kg, and average daily load."),
        ("Baseline Savings", "Verified Working", "calc.py:70-72, pages.jsx:26-30", "Correctly handles negative savings; displays 'Net vs baseline' and excess energy when usage exceeds baseline."),
        ("Stability Score", "Verified Working", "calc.py:112-120, pages.jsx:33", "Correctly labeled as a heuristic stability index (0–100), not a building efficiency certificate. Clamped to [0, 100]."),
        ("Period Resampling", "Verified Working", "calc.py:40-47, pages.jsx:11-12", "Supports Daily (last 30 days), Weekly (week ending Sunday), and Monthly aggregations."),
        ("Category Breakdown", "Verified Working", "calc.py:82-85, pages.jsx:36-39", "Shows category kWh and share percentages for Cooling, Lighting, Appliances, Office Equipment."),
        ("Manual Record CRUD", "Verified Working", "app.py:38-69, pages.jsx:41-75", "Supports Create, Read, Update, and Delete with input validation and instant state refresh."),
        ("CSV Ingestion Engine", "Verified Working", "app.py:71-89, calc.py:10-32", "Ingests CSV files, reports added and skipped rows, skips invalid entries, validates dates and kWh."),
        ("CSV Template Download", "Verified Working", "pages.jsx:51-52", "Generates wattwise_template.csv directly in browser with valid pre-populated headers."),
        ("Demo Data Controls", "Verified Working", "app.py:142-148, store.py:11-25", "One-click reset to 45 days of demo data (with weekend patterns and AC spikes) or complete purge."),
        ("High-Consumption Days", "Verified Working", "calc.py:73-76, pages.jsx:81", "Identifies days where daily consumption exceeds mean + 1.5 standard deviations (requires >= 7 days)."),
        ("Weekday vs Weekend", "Verified Working", "calc.py:49-55, pages.jsx:82", "Computes weekday vs. weekend averages when at least 4 weekdays and 4 weekend days exist."),
        ("Spike Anomaly Alert", "Verified Working", "calc.py:77-79, pages.jsx:83", "Flags days where consumption exceeds 1.30x trailing 7-day rolling mean."),
        ("Explainable Rules", "Verified Working", "calc.py:86-111, pages.jsx:13-15", "Transparently states trigger rule, explanation, and estimated monthly kWh and ₹ savings."),
        ("ML Forecasting", "Verified Working", "forecast.py:1-42, pages.jsx:91-114", "Features 7d, 14d, and 30d presets; compares 7-day mean baseline vs Ridge regression; reports test MAE."),
        ("Cost & Carbon Calc", "Verified Working", "calc.py:139-144, pages.jsx:116-129", "What-If scenario simulator calculating savings across custom percentage reduction targets."),
        ("Solar PV & Payback", "Verified Working", "calc.py:125-138, pages.jsx:131-157", "Computes daily/monthly generation, rooftop limits, avoided emissions, turnkey cost, and simple payback."),
        ("CSV Reports Export", "Verified Working", "app.py:150-160, pages.jsx:159-170", "Downloads raw itemized records (wattwise_records.csv) or periodic summary reports (wattwise_report.csv).")
    ]

    for idx, (f_name, f_status, f_evid, f_notes) in enumerate(verification_data):
        row = table_feat.rows[idx + 1]
        for c_idx, val in enumerate([f_name, f_status, f_evid, f_notes]):
            cell = row.cells[c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            if c_idx == 0:
                r.font.bold = True
            elif c_idx == 1:
                r.font.bold = True
                r.font.color.rgb = PRIMARY
            bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 35, 35, 50, 50)

    # Section 5: Mathematical Formulas
    add_heading_1("5. Mathematical Formulas & Engineering Equations")

    formulas = [
        ("Electricity Cost (₹)", "Cost = Sum(Consumption_i * Tariff_i)"),
        ("Carbon Footprint (kg CO2e)", "Emissions = Total_kWh * Emission_Factor (Default: 0.82 kg CO2e/kWh, configurable)"),
        ("Baseline Savings (kWh)", "Saved_kWh = (Baseline_Daily_Avg * Total_Days) - Total_kWh (Negative values denote excess)"),
        ("Anomaly Spike Threshold", "Spike_Condition = Daily_kWh > 1.30 * Rolling_7_Day_Mean"),
        ("Stability Score (0–100)", "Score = 100 - (30 * min(CV/0.4, 1)) - (40 * min(max(WoW_Change, 0)/30, 1)) - (5 * min(Anomalies, 6))"),
        ("Solar Capacity Sizing (kW)", "Capacity = min(Rooftop_Area / 6.0 m2/kW, Daily_kWh / (Peak_Sun_Hours * Performance_Factor))"),
        ("Solar Daily Output (kWh)", "Daily_Generation = System_Capacity_kW * Sun_Hours * Performance_Ratio (0.75)"),
        ("Solar Simple Payback", "Payback_Years = (System_Capacity_kW * 50,000 ₹/kW) / Annual_Cost_Savings_INR")
    ]

    table_form = doc.add_table(rows=len(formulas)+1, cols=2)
    table_form.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table_form.rows[0]
    hdr.cells[0].paragraphs[0].add_run("Metric / Equation").font.bold = True
    hdr.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    hdr.cells[1].paragraphs[0].add_run("Mathematical Definition").font.bold = True
    hdr.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_background(hdr.cells[0], "166534")
    set_cell_background(hdr.cells[1], "166534")

    for idx, (name, eq) in enumerate(formulas):
        r = table_form.rows[idx+1]
        c0, c1 = r.cells[0], r.cells[1]
        c0.paragraphs[0].add_run(name).font.bold = True
        c1.paragraphs[0].add_run(eq)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 60, 60)

    # Section 6: Machine Learning Methodology
    add_heading_1("6. Machine Learning Forecasting Methodology")
    doc.add_paragraph(
        "Forecasting uses an L2 regularized Ridge Regression pipeline (alpha=1.0) with timeline trend and "
        "one-hot day-of-week feature encoding. To prevent future data leakage, training uses a strict chronological "
        "split where the final 20% of history is reserved as hold-out test data."
    )
    doc.add_paragraph(
        "The model is compared against a naive 7-day moving average baseline on identical test observations using Mean Absolute Error (MAE). "
        "If the regression model does not beat the baseline on held-out data, the system automatically falls back to the labeled baseline forecast."
    )

    # Section 7: Verification & Quality Assurance
    add_heading_1("7. Quality Assurance & Test Evidence")
    doc.add_paragraph("All automated tests and production builds were executed and verified during this evaluation:")
    
    test_bullets = [
        "Backend Unit Tests (pytest test_wattwise.py): 10 passed out of 10 tests (100% pass rate in 11.69s).",
        "Test Coverage: Tests validate percentage changes, data validation, solar physical math, scenario savings, CRUD operations, CSV ingestion, ML forecasting fallbacks, negative baseline savings handling, and solar payback limits.",
        "Frontend Production Compilation (vite build): Successfully compiled 2,382 modules in 24.08s with zero runtime or syntax errors.",
        "CORS Security: Configured to explicitly allow localhost development ports (http://localhost:5173, http://127.0.0.1:5173)."
    ]
    for b in test_bullets:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(b)

    # Section 8: Limitations & Future Scope
    add_heading_1("8. Limitations & Unresolved Scope")
    limitations = [
        "No Live IoT Telemetry: The system processes user-entered records and CSV uploads; it is not physically connected to smart meters or CT clamps.",
        "Unverified Default Emissions Factor: The default 0.82 kg CO2e/kWh is a national aggregate placeholder. Verified regional factors (such as CEA Baseline v19) can be configured via settings.",
        "Linear Solar Sizing: Sizing is based on geometric rooftop area and average peak sun hours; it does not perform 3D ray-tracing or hourly micro-climate simulations.",
        "File-Based JSON Storage: data.json provides lightweight zero-config persistence for hackathon deployment; enterprise scaling requires migration to PostgreSQL."
    ]
    for lim in limitations:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(lim)

    # Section 9: Presentation Walkthrough
    add_heading_1("9. InnoTech Hackathon Presentation Walkthrough")
    walkthrough = [
        ("Step 1: Introduction", "Introduce WattWise AI as an accessible, software-driven energy intelligence platform for campuses and commercial buildings."),
        ("Step 2: Live Dashboard", "Show real-time calculated KPI cards, 0–100 consumption stability index, and daily/weekly/monthly Recharts views."),
        ("Step 3: Data Ingestion & Demo Mode", "Show manual record addition with validation, one-click demo data reset, and test upload using sample_consumption.csv."),
        ("Step 4: AI & ML Engine", "Demonstrate transparent recommendations (with explicit kWh and ₹ savings) and run the ML forecast showing MAE model vs. baseline accuracy comparison across 7d, 14d, and 30d."),
        ("Step 5: Renewable Solar Planner", "Enter campus rooftop dimensions to showcase instant solar system sizing, annual bill savings, and payback years."),
        ("Step 6: Data Export", "Demonstrate CSV export functionality for reporting and enterprise integration.")
    ]
    for title, desc in walkthrough:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(f"• {title}: ")
        r_title.font.bold = True
        r_title.font.color.rgb = PRIMARY
        p.add_run(desc)

    # Section 10: References
    add_heading_1("10. Verified References")
    refs = [
        "Bureau of Energy Efficiency (BEE), Ministry of Power, India: Energy Conservation Guidelines for Air Conditioning in Commercial Buildings (24°C setpoint guideline).",
        "Central Electricity Authority (CEA), Government of India: CO2 Baseline Database for the Indian Power Sector, Version 19.",
        "National Renewable Energy Laboratory (NREL): PVWatts Photovoltaic Sizing and Performance Ratio Guidelines (0.75-0.80 derate ratio).",
        "Scikit-learn Documentation: L2 Regularized Ridge Regression (sklearn.linear_model.Ridge)."
    ]
    for ref in refs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(ref)

    # Save Document
    output_path = r"d:\websites\hackthon\WattWise_AI_Technical_Report_Final.docx"
    doc.save(output_path)
    print(f"Report saved successfully to {output_path}")

if __name__ == "__main__":
    create_report()
